"""Report Engine: Generates multi-format audit reports (HTML, PDF, DOCX, JSON)."""

from __future__ import annotations

import io
import json
from typing import Any

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class ReportEngine:
    """Generates official audit reports in various formats."""

    def __init__(
        self,
        summary: dict[str, Any],
        findings: list[dict[str, Any]],
        document_info: dict[str, Any],
        checklist: dict[str, list[dict[str, Any]]],
        category_scores: dict[str, float] | None = None,
        overall_score: float = 85.0
    ):
        self.summary = summary
        self.findings = findings
        self.doc_info = document_info
        self.checklist = checklist
        if isinstance(category_scores, (int, float)):
            self.overall_score = float(category_scores)
            self.category_scores = {}
        else:
            self.category_scores = category_scores or {}
            self.overall_score = float(overall_score)

    def to_json(self) -> str:
        """Exports audit report as structured JSON."""
        data = {
            "title": "LAPORAN AUDIT KEPATUHAN & KONSISTENSI TESIS MIKOM 2026",
            "document": self.doc_info,
            "overall_score": self.overall_score,
            "category_scores": self.category_scores,
            "summary": self.summary,
            "findings_count": len(self.findings),
            "findings": self.findings,
            "revision_checklist": self.checklist
        }
        return json.dumps(data, indent=2, ensure_ascii=False)

    def to_html(self) -> str:
        """Exports executive HTML audit report."""
        rows = ""
        for f in self.findings:
            status_color = "#DC2626" if f.get("status") == "ERROR" else "#D97706"
            rows += f"""
            <tr style="border-bottom: 1px solid #E2E8F0;">
                <td style="padding: 10px; font-weight: bold; color: {status_color};">{f.get('status')}</td>
                <td style="padding: 10px; font-family: monospace;">{f.get('rule_id')}</td>
                <td style="padding: 10px;">{f.get('category')}</td>
                <td style="padding: 10px;"><b>Hal. {f.get('page') or '-'}</b>: {f.get('description')}</td>
                <td style="padding: 10px; font-size: 0.9em; background: #F8FAFC;">{f.get('actual')}</td>
                <td style="padding: 10px; font-size: 0.9em;">{f.get('recommendation')}</td>
            </tr>
            """

        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>Laporan Audit Tesis MIKOM 2026</title>
            <style>
                body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; margin: 40px; color: #1E293B; }}
                h1 {{ color: #0F172A; border-bottom: 2px solid #3B82F6; padding-bottom: 10px; }}
                .badge {{ padding: 6px 12px; border-radius: 6px; font-weight: bold; color: white; display: inline-block; }}
                .score {{ font-size: 32px; font-weight: bold; color: #2563EB; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
                th {{ background: #F1F5F9; padding: 12px; text-align: left; font-size: 0.85em; text-transform: uppercase; }}
            </style>
        </head>
        <body>
            <h1>Laporan Audit Kepatuhan & Konsistensi Tesis MIKOM 2026</h1>
            <p><b>Dokumen:</b> {self.doc_info.get('filename')} | <b>Halaman:</b> {self.doc_info.get('page_count')} | <b>Total Temuan:</b> {len(self.findings)}</p>
            <div style="background: #EFF6FF; border: 1px solid #BFDBFE; border-radius: 8px; padding: 16px; margin: 20px 0;">
                <span class="score">{self.overall_score} / 100</span>
                <p style="margin: 5px 0 0 0; color: #1E40AF;">Indeks Kepatuhan & Konsistensi Ilmiah Tesis</p>
            </div>
            <table>
                <thead>
                    <tr>
                        <th>Status</th>
                        <th>ID Aturan</th>
                        <th>Kategori</th>
                        <th>Deskripsi & Lokasi</th>
                        <th>Kondisi Terdeteksi</th>
                        <th>Rekomendasi Revisi</th>
                    </tr>
                </thead>
                <tbody>
                    {rows}
                </tbody>
            </table>
        </body>
        </html>
        """
        return html

    def to_docx_bytes(self) -> io.BytesIO:
        """Exports audit report as a DOCX document."""
        doc = Document()
        doc.add_heading("LAPORAN AUDIT TESIS MIKOM 2026", level=1)
        doc.add_paragraph(f"Dokumen: {self.doc_info.get('filename')} | Skor: {self.overall_score}/100")

        # Summary counts
        p = doc.add_paragraph()
        p.add_run(f"PASS: {self.summary.get('PASS', 0)}  |  ERROR: {self.summary.get('ERROR', 0)}  |  WARNING: {self.summary.get('WARNING', 0)}  |  UNVERIFIED: {self.summary.get('UNVERIFIED', 0)}").bold = True

        doc.add_heading("Rincian Temuan Audit", level=2)
        for f in self.findings:
            item_p = doc.add_paragraph()
            run_status = item_p.add_run(f"[{f.get('status')}] {f.get('rule_id')}: {f.get('description')}\n")
            run_status.bold = True
            item_p.add_run(f"Lokasi: Halaman {f.get('page') or '-'} | Bagian: {f.get('section') or '-'}\n")
            item_p.add_run(f"Aktual: {f.get('actual')}\n")
            item_p.add_run(f"Ekspektasi: {f.get('expected')}\n")
            item_p.add_run(f"Rekomendasi: {f.get('recommendation')}\n")

        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)
        return bio

    def to_pdf_bytes(self) -> io.BytesIO:
        """Exports audit report as PDF using ReportLab."""
        bio = io.BytesIO()
        pdf = SimpleDocTemplate(bio, pagesize=A4)
        styles = getSampleStyleSheet()

        elements = []
        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            textColor=colors.HexColor("#1E3A8A"),
            spaceAfter=12
        )
        elements.append(Paragraph("LAPORAN AUDIT TESIS MIKOM 2026", title_style))
        elements.append(Paragraph(f"Dokumen: {self.doc_info.get('filename')} | Skor Kepatuhan: {self.overall_score}/100", styles["Normal"]))
        elements.append(Spacer(1, 15))

        # Table of findings
        table_data = [["Status", "ID Aturan", "Halaman", "Deskripsi", "Rekomendasi"]]
        for f in self.findings[:25]:  # Top 25 findings
            table_data.append([
                str(f.get("status", "")),
                str(f.get("rule_id", "")),
                str(f.get("page") or "-"),
                str(f.get("description", ""))[:45],
                str(f.get("recommendation", ""))[:45]
            ])

        t = Table(table_data, colWidths=[50, 75, 45, 170, 170])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elements.append(t)
        pdf.build(elements)
        bio.seek(0)
        return bio


def generate_plagiarism_report(report: Any, segments: list[dict[str, Any]]) -> bytes:
    """Menghasilkan laporan PDF ringkasan plagiarisme & konten AI."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Laporan Plagiarisme & Konten AI", styles["Title"]))
    story.append(Spacer(1, 12))

    # Ringkasan skor
    summary = [
        ["Metrik", "Nilai"],
        ["Skor Plagiarisme", f"{report.overall_plagiarism_score*100:.1f}%"],
        ["Skor Konten AI", f"{report.overall_ai_score*100:.1f}%"],
        ["Status", report.verdict.upper()],
        ["Segmen Diperiksa", str(len(segments))],
    ]
    t = Table(summary, colWidths=[200, 200])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(t)
    story.append(Spacer(1, 20))

    # Segmen mencurigakan
    story.append(Paragraph("<b>Segmen dengan Kemiripan Tinggi</b>", styles["Heading2"]))
    suspicious = sorted(
        [s for s in segments if s.get("score", 0) >= 0.40],
        key=lambda s: s.get("score", 0),
        reverse=True,
    )[:30]
    for i, s in enumerate(suspicious, 1):
        story.append(Paragraph(
            f"<b>{i}. Skor {s['score']*100:.0f}%</b> ({s.get('checker', '-')})",
            styles["Normal"]
        ))
        story.append(Paragraph(f"<i>{s['text'][:300]}</i>", styles["Normal"]))
        if s.get("source"):
            story.append(Paragraph(f"Sumber: {s['source'][:150]}", styles["Normal"]))
        story.append(Spacer(1, 6))

    doc.build(story)
    return buf.getvalue()


def generate_citation_report(findings: list[Any]) -> bytes:
    """Menghasilkan laporan PDF ringkasan audit sitasi Scopus & SINTA."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Laporan Audit Verifikasi Sitasi (Scopus & SINTA)", styles["Title"]))
    story.append(Spacer(1, 12))

    status_counts: dict[str, int] = {}
    for f in findings:
        status = getattr(f, "status", "unknown")
        status_counts[status] = status_counts.get(status, 0) + 1

    summary = [["Status Verifikasi", "Jumlah Sitasi"]]
    for st_name, count in status_counts.items():
        summary.append([st_name.upper(), str(count)])

    t = Table(summary, colWidths=[200, 200])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(t)
    story.append(Spacer(1, 20))

    story.append(Paragraph("<b>Daftar Temuan Sitasi</b>", styles["Heading2"]))
    for i, f in enumerate(findings[:30], 1):
        txt = getattr(f, "citation_text", str(f))
        status = getattr(f, "status", "info")
        ev = getattr(f, "evidence_ref", "")
        line = f"<b>{i}. [{status.upper()}]</b> {txt[:200]}"
        if ev:
            line += f" ({ev})"
        story.append(Paragraph(line, styles["Normal"]))
        story.append(Spacer(1, 4))

    doc.build(story)
    return buf.getvalue()


def generate_full_report(
    meta: dict[str, Any],
    format_findings: list[Any],
    structure_findings: list[Any],
    citation_findings: list[Any],
    plagiarism_report: Any,
    sections: list[str],
    title: str,
    author: str,
    language: str,
    include_evidence: bool,
) -> bytes:
    """Menghasilkan laporan PDF komprehensif seluruh modul audit."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    # Cover
    story.append(Paragraph(title, styles["Title"]))
    if author:
        story.append(Paragraph(f"Disusun oleh: {author}", styles["Normal"]))
    story.append(Paragraph(f"Dokumen: {meta.get('filename', '-')}", styles["Normal"]))
    story.append(PageBreak())

    # Ringkasan Eksekutif
    if "Ringkasan Eksekutif" in sections:
        story.append(Paragraph("Ringkasan Eksekutif", styles["Heading1"]))
        plag_score = getattr(plagiarism_report, "overall_plagiarism_score", 0.0) * 100
        ai_score = getattr(plagiarism_report, "overall_ai_score", 0.0) * 100
        story.append(Paragraph(
            f"Total temuan: format={len(format_findings)}, "
            f"struktur={len(structure_findings)}, "
            f"sitasi={len(citation_findings)}. "
            f"Skor plagiarisme={plag_score:.0f}%, "
            f"AI={ai_score:.0f}%.",
            styles["Normal"]
        ))
        story.append(PageBreak())

    # Format
    if "Format" in sections and format_findings:
        story.append(Paragraph("Temuan Format", styles["Heading1"]))
        for f in format_findings:
            sev = getattr(f, "severity", "INFO")
            msg = getattr(f, "message", getattr(f, "description", str(f)))
            story.append(Paragraph(f"• [{sev}] {msg}", styles["Normal"]))
        story.append(PageBreak())

    # Struktur
    if "Struktur" in sections and structure_findings:
        story.append(Paragraph("Temuan Struktur", styles["Heading1"]))
        for f in structure_findings:
            sev = getattr(f, "severity", "INFO")
            msg = getattr(f, "message", getattr(f, "description", str(f)))
            story.append(Paragraph(f"• [{sev}] {msg}", styles["Normal"]))
        story.append(PageBreak())

    # Sitasi
    if "Sitasi" in sections and citation_findings:
        story.append(Paragraph("Verifikasi Sitasi", styles["Heading1"]))
        for f in citation_findings:
            st_text = getattr(f, "status", "info")
            cite_txt = getattr(f, "citation_text", str(f))[:200]
            line = f"• [{st_text}] {cite_txt}"
            if include_evidence and hasattr(f, "evidence_ref") and f.evidence_ref:
                line += f" — {f.evidence_ref}"
            story.append(Paragraph(line, styles["Normal"]))
        story.append(PageBreak())

    # Plagiarisme
    if "Plagiarisme" in sections:
        story.append(Paragraph("Plagiarisme & Konten AI", styles["Heading1"]))
        plag_score = getattr(plagiarism_report, "overall_plagiarism_score", 0.0) * 100
        ai_score = getattr(plagiarism_report, "overall_ai_score", 0.0) * 100
        verdict = getattr(plagiarism_report, "verdict", "clean")
        story.append(Paragraph(
            f"Skor plagiarisme: {plag_score:.1f}%<br/>"
            f"Skor AI: {ai_score:.1f}%<br/>"
            f"Status: {verdict.upper()}",
            styles["Normal"]
        ))

    doc.build(story)
    return buf.getvalue()


def generate_full_report_docx(
    meta: dict[str, Any],
    format_findings: list[Any],
    structure_findings: list[Any],
    citation_findings: list[Any],
    plagiarism_report: Any,
    sections: list[str],
    title: str,
    author: str,
) -> bytes:
    """Menghasilkan laporan DOCX komprehensif seluruh modul audit."""
    doc = Document()

    # Cover
    doc.add_heading(title, level=0)
    if author:
        doc.add_paragraph(f"Disusun oleh: {author}")
    doc.add_paragraph(f"Dokumen: {meta.get('filename', '-')}")
    doc.add_page_break()

    if "Ringkasan Eksekutif" in sections:
        doc.add_heading("Ringkasan Eksekutif", level=1)
        doc.add_paragraph(
            f"Total temuan: format={len(format_findings)}, "
            f"struktur={len(structure_findings)}, "
            f"sitasi={len(citation_findings)}."
        )
        doc.add_page_break()

    if "Format" in sections:
        doc.add_heading("Temuan Format", level=1)
        for f in format_findings:
            sev = getattr(f, "severity", "INFO")
            msg = getattr(f, "message", getattr(f, "description", str(f)))
            doc.add_paragraph(f"[{sev}] {msg}", style="List Bullet")
        doc.add_page_break()

    if "Struktur" in sections:
        doc.add_heading("Temuan Struktur", level=1)
        for f in structure_findings:
            sev = getattr(f, "severity", "INFO")
            msg = getattr(f, "message", getattr(f, "description", str(f)))
            doc.add_paragraph(f"[{sev}] {msg}", style="List Bullet")
        doc.add_page_break()

    if "Sitasi" in sections:
        doc.add_heading("Verifikasi Sitasi", level=1)
        for f in citation_findings:
            st_text = getattr(f, "status", "info")
            cite_txt = getattr(f, "citation_text", str(f))[:200]
            ev = getattr(f, "evidence_ref", "") or ""
            line = f"[{st_text}] {cite_txt}"
            if ev:
                line += f" — {ev}"
            doc.add_paragraph(line, style="List Bullet")
        doc.add_page_break()

    if "Plagiarisme" in sections:
        doc.add_heading("Plagiarisme & Konten AI", level=1)
        plag_score = getattr(plagiarism_report, "overall_plagiarism_score", 0.0) * 100
        ai_score = getattr(plagiarism_report, "overall_ai_score", 0.0) * 100
        doc.add_paragraph(f"Skor plagiarisme: {plag_score:.1f}%")
        doc.add_paragraph(f"Skor AI: {ai_score:.1f}%")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

