from docx import Document

from engines.docx_parser import parse_docx


def test_parse_docx(tmp_path):
    path = tmp_path / "sample.docx"

    doc = Document()
    doc.add_heading("BAB I PENDAHULUAN", level=1)
    doc.add_paragraph("Latar belakang penelitian.")
    doc.save(path)

    result = parse_docx(str(path))

    assert result.file_type == "docx"
    assert "Latar belakang penelitian." in result.get_text()
    assert result.page_count is None
    assert len(result.paragraphs) >= 2
