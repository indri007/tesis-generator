"""tests/unit/test_premium_academic_tools.py - Unit tests for premium docx, plot, and palette."""
import io
import pytest
from docx import Document
from utils.premium_docx import apply_premium_typography, set_booktabs_style, generate_sample_premium_docx
from utils.premium_plot import create_premium_chart, MUTED_PALETTE
from utils.palette_generator import generate_premium_palette


def test_premium_typography():
    doc = Document()
    apply_premium_typography(doc)
    style = doc.styles['Normal']
    assert style.font.name == 'Times New Roman'
    assert style.paragraph_format.widow_control is True


def test_booktabs_style():
    doc = Document()
    apply_premium_typography(doc)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Header 1"
    table.cell(1, 0).text = "Data 1"
    set_booktabs_style(table)
    
    buf = io.BytesIO()
    doc.save(buf)
    assert len(buf.getvalue()) > 0


def test_generate_sample_premium_docx():
    buf = generate_sample_premium_docx(title="BAB I UJI COBA")
    assert isinstance(buf, io.BytesIO)
    content = buf.getvalue()
    assert len(content) > 1000
    # Verifikasi dapat dibuka kembali sebagai docx yang valid
    doc = Document(io.BytesIO(content))
    assert len(doc.tables) == 1
    assert len(doc.paragraphs) >= 2


def test_create_premium_chart_png_and_svg():
    x = [2022, 2023, 2024, 2025, 2026]
    y = [45, 62, 78, 85, 94]
    
    png_buf = create_premium_chart(
        x_data=x,
        y_data=y,
        title="Tren Keterlibatan Publik",
        xlabel="Tahun",
        ylabel="Skor (%)",
        annotation_text="Puncak Program MBG",
        export_fmt="png"
    )
    assert len(png_buf.getvalue()) > 1000
    
    svg_buf = create_premium_chart(
        x_data=x,
        y_data=y,
        title="Tren Keterlibatan Publik",
        xlabel="Tahun",
        ylabel="Skor (%)",
        annotation_text="Puncak Program MBG",
        export_fmt="svg"
    )
    svg_text = svg_buf.getvalue().decode('utf-8', errors='ignore')
    assert "<svg" in svg_text
    assert "Tren Keterlibatan Publik" in svg_text


def test_generate_premium_palette():
    palette = generate_premium_palette()
    assert "Primary (Utama)" in palette
    assert "Secondary (Sekunder)" in palette
    assert "Tertiary (Netral/Latar)" in palette
    for k, v in palette.items():
        assert v.startswith("#")
