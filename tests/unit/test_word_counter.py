"""Unit tests for word counter excluding figures, captions, and tables."""

from pathlib import Path
from docx import Document
import pytest

from utils.word_counter import count_words_excl_figures, evaluate_word_count, WORD_MIN, WORD_MAX, MAX_PAGES


def test_max_pages_none_and_word_targets():
    """Verify MAX_PAGES is None and word boundaries are 40k and 60k."""
    assert MAX_PAGES is None
    assert WORD_MIN == 40000
    assert WORD_MAX == 60000


def test_evaluate_word_count_statuses():
    """Verify evaluation statuses according to boundaries."""
    status_err, _, _ = evaluate_word_count(35000)
    assert status_err == "ERROR"

    status_pass, _, _ = evaluate_word_count(45000)
    assert status_pass == "PASS"

    status_warn, _, _ = evaluate_word_count(65000)
    assert status_warn == "WARNING"


def test_count_words_excl_figures_docx(tmp_path: Path):
    """Verify words before BAB I and after DAFTAR PUSTAKA as well as captions are ignored."""
    docx_file = tmp_path / "test_manuscript.docx"
    doc = Document()

    # Front matter (should be ignored)
    doc.add_paragraph("KATA PENGANTAR Pembuka naskah tesis.")
    doc.add_paragraph("ABSTRAK Penelitian ini mengkaji komunikasi pemasaran digital.")

    # Bab I (included)
    p_b1 = doc.add_paragraph("BAB I PENDAHULUAN")
    p_b1.style = "Heading 1"
    doc.add_paragraph("Latar belakang riset ini sangat penting dalam era transformasi digital.")

    # Caption (should be excluded)
    doc.add_paragraph("Gambar 1.1 Diagram Alur Penelitian Kerangka Konseptual")
    doc.add_paragraph("Tabel 1.2 Data Responden Survei Lapangan")

    # Bab II-V text (included)
    doc.add_paragraph("Kajian teori Uses and Gratifications dan Media Richness membuktikan relevansi.")

    # Daftar Pustaka (break loop)
    p_dp = doc.add_paragraph("DAFTAR PUSTAKA")
    p_dp.style = "Heading 1"
    doc.add_paragraph("Kartika, I. (2026). Jurnal Komunikasi.")

    doc.save(str(docx_file))

    total_words = count_words_excl_figures(str(docx_file))
    # Should only count:
    # "BAB I PENDAHULUAN" (3 words)
    # "Latar belakang riset ini sangat penting dalam era transformasi digital." (10 words)
    # "Kajian teori Uses and Gratifications dan Media Richness membuktikan relevansi." (10 words)
    # Total = 23 words
    assert total_words == 23
