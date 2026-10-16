"""Unit tests for DocxAutoFixer engine in accordance with MIKOM 2026 guidelines."""

from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from engines.revision_engine import DocxAutoFixer


def test_docx_auto_fixer_applies_mikom_2026_formatting(tmp_path: Path):
    # 1. Create a dummy docx with default/wrong formatting
    in_file = tmp_path / "thesis_draft_wrong_format.docx"
    doc = Document()
    p1 = doc.add_paragraph("Ini adalah contoh paragraf pertama dengan format bawaan.")
    p2 = doc.add_paragraph("Ini adalah contoh paragraf kedua yang perlu dirapikan.")
    doc.save(str(in_file))

    orig_bytes = in_file.read_bytes()

    # 2. Apply DocxAutoFixer
    out_file = tmp_path / "thesis_draft_mikom_fixed.docx"
    fixer = DocxAutoFixer(str(in_file))
    log = fixer.fix_all(str(out_file))

    # 3. Assert non-destructive source file integrity
    assert in_file.read_bytes() == orig_bytes, "Source file must not be modified!"

    # 4. Assert output file exists and is valid
    assert out_file.exists()
    assert len(log) >= 2

    # 5. Inspect modified document properties
    fixed_doc = Document(str(out_file))
    for section in fixed_doc.sections:
        assert round(section.top_margin.cm, 1) == 4.0
        assert round(section.left_margin.cm, 1) == 4.0
        assert round(section.bottom_margin.cm, 1) == 3.0
        assert round(section.right_margin.cm, 1) == 3.0

    for p in fixed_doc.paragraphs:
        assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert p.paragraph_format.line_spacing == 2.0
        for r in p.runs:
            assert r.font.name == "Times New Roman"
            assert r.font.size == Pt(12)


def test_docx_auto_fixer_preserves_headings_and_tables(tmp_path: Path):
    """Verify that headings keep their alignment and tables are formatted safely."""
    in_file = tmp_path / "thesis_with_table.docx"
    doc = Document()
    h = doc.add_paragraph("BAB I PENDAHULUAN")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Variabel"
    tbl.cell(0, 1).text = "Indikator"
    doc.save(str(in_file))

    out_file = tmp_path / "thesis_table_fixed.docx"
    fixer = DocxAutoFixer(str(in_file))
    log = fixer.fix_all(str(out_file))

    assert out_file.exists()
    actions = [item["action"] for item in log]
    assert "SET_MARGINS_4_4_3_3_CM" in actions
    assert "SET_FONT_AND_SPACING" in actions
    assert "FORMAT_TABLE_CELLS" in actions

    fixed_doc = Document(str(out_file))
    # Heading kept center alignment
    assert fixed_doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_docx_auto_fixer_error_handling(tmp_path: Path):
    """Verify proper exceptions on missing file or corrupt file."""
    import pytest

    # Non-existent file
    with pytest.raises(FileNotFoundError):
        DocxAutoFixer(str(tmp_path / "does_not_exist.docx"))

    # Corrupt file
    corrupt_file = tmp_path / "corrupt.docx"
    corrupt_file.write_bytes(b"Corrupt binary data that is not a zip or docx")
    fixer = DocxAutoFixer(str(corrupt_file))
    with pytest.raises(ValueError, match="rusak atau tidak dapat diuraikan"):
        fixer.fix_all(str(tmp_path / "never_created.docx"))

