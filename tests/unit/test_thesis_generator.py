"""Unit tests for ThesisGenerator: page allocation, chapter generation, and DOCX styling."""

import io

from docx import Document

from core.generator.thesis_generator import ThesisGenerator


def test_page_budget_calculation():
    """Verify page allocation adheres to 15%, 25%, 15%, 35%, 10% proportions and 175 page limit."""
    budget = ThesisGenerator.calculate_page_budget(target_pages=100)
    assert budget["BAB I"]["target_pages"] == 15
    assert budget["BAB II"]["target_pages"] == 25
    assert budget["BAB III"]["target_pages"] == 15
    assert budget["BAB IV"]["target_pages"] == 35
    assert budget["BAB V"]["target_pages"] == 10

    # Over 175 pages hard limit capped to 175
    budget_capped = ThesisGenerator.calculate_page_budget(target_pages=200)
    total_pages = sum(b["target_pages"] for b in budget_capped.values())
    assert total_pages == 175


def test_chapter_generation_consistency():
    """Verify that variables X and Y are consistent across all generated chapters."""
    gen = ThesisGenerator()
    params = {
        "title": "Pengaruh Komunikasi Digital terhadap Loyalitas Audiens",
        "var_x": "Kualitas Pesan Digital",
        "var_y": "Loyalitas Audiens",
        "theory": "Media Richness Theory",
        "sample_n": 120,
        "formula_a": 2.5,
        "formula_b": 0.4,
        "eval_x": 10.0,
        "p_value": 0.032,
    }

    c1 = gen.generate_chapter_1(params)
    c2 = gen.generate_chapter_2(params)
    c3 = gen.generate_chapter_3(params)
    c4 = gen.generate_chapter_4(params)
    c5 = gen.generate_chapter_5(params)

    # Check variable consistency across chapters
    assert "Kualitas Pesan Digital" in c1
    assert "Loyalitas Audiens" in c1
    assert "Kualitas Pesan Digital" in c2
    assert "H1:" in c2
    assert "120 responden" in c3
    assert "120 responden" in c4
    assert "Y = 2,5 + 0,4X" in c4
    assert "Y = 6,5" in c4  # 2.5 + 0.4(10) = 6.5
    assert "Kualitas Pesan Digital" in c5


def test_compile_docx_formatting():
    """Verify compiled DOCX document has correct margins (4-4-3-3 cm) and Times New Roman font."""
    gen = ThesisGenerator()
    params = {
        "title": "Pengaruh Komunikasi Digital terhadap Loyalitas Audiens",
        "author": "Indri Kartika",
        "nim": "2401234567",
        "var_x": "Kualitas Pesan Digital",
        "var_y": "Loyalitas Audiens",
    }
    docx_stream = gen.compile_full_thesis_docx(params)
    assert isinstance(docx_stream, io.BytesIO)

    # Open document and inspect properties
    doc = Document(docx_stream)
    section = doc.sections[0]

    # Margins in cm (rounded to 1 decimal place)
    assert round(section.top_margin.cm, 1) == 4.0
    assert round(section.left_margin.cm, 1) == 4.0
    assert round(section.bottom_margin.cm, 1) == 3.0
    assert round(section.right_margin.cm, 1) == 3.0

    # Ensure paragraphs exist and have font
    assert len(doc.paragraphs) > 20
    first_run = doc.paragraphs[0].runs[0]
    assert first_run.font.name == "Times New Roman"
    assert first_run.font.size.pt == 12.0
