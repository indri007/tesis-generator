"""Unit tests for UnifiedDocumentExtractor, OCREngine, and DocumentConverter."""

from pathlib import Path
import pytest
from docx import Document
from PIL import Image, ImageDraw

from core.document_processing.unified_extractor import (
    UnifiedDocumentExtractor,
    DocumentValidationError,
)
from core.ocr.ocr_engine import OCREngine
from core.conversion.document_converter import DocumentConverter

try:
    import pymupdf as fitz
except ImportError:
    import fitz


@pytest.fixture
def extractor() -> UnifiedDocumentExtractor:
    return UnifiedDocumentExtractor(max_size_mb=10)


def test_pdf_digital_text_extraction(extractor: UnifiedDocumentExtractor, tmp_path: Path):
    """Test 1: PDF digital yang memiliki teks berhasil diekstraksi."""
    pdf_file = tmp_path / "digital_sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 100), "Analisis Komunikasi Organisasi di Era Transformasi Digital.")
    doc.save(str(pdf_file))
    doc.close()

    text = extractor.extract_pdf_text(pdf_file)
    assert "Analisis Komunikasi Organisasi" in text


def test_docx_structure_and_table_extraction(extractor: UnifiedDocumentExtractor, tmp_path: Path):
    """Test 2: DOCX dengan heading, paragraf, dan tabel berhasil diekstraksi."""
    docx_file = tmp_path / "sample_academic.docx"
    doc = Document()
    doc.add_heading("BAB I PENDAHULUAN", level=1)
    doc.add_paragraph("Latar belakang penelitian ini difokuskan pada komunikasi publik.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Indikator"
    table.cell(0, 1).text = "Skor"
    table.cell(1, 0).text = "Transparansi"
    table.cell(1, 1).text = "85%"
    doc.save(str(docx_file))

    text = extractor.extract_docx_text(docx_file)
    assert "BAB I PENDAHULUAN" in text
    assert "komunikasi publik" in text
    assert "Transparansi | 85%" in text


def test_empty_or_corrupted_file_validation(extractor: UnifiedDocumentExtractor, tmp_path: Path):
    """Test 3: File kosong atau rusak memicu DocumentValidationError."""
    empty_file = tmp_path / "empty.docx"
    empty_file.touch()

    with pytest.raises(DocumentValidationError, match="Berkas kosong"):
        extractor.validate_file(empty_file)

    corrupt_pdf = tmp_path / "corrupt.pdf"
    corrupt_pdf.write_bytes(b"Bukan PDF yang valid sama sekali")

    with pytest.raises(DocumentValidationError, match="Signature berkas tidak valid"):
        extractor.validate_file(corrupt_pdf)


def test_mismatched_extension_validation(extractor: UnifiedDocumentExtractor, tmp_path: Path):
    """Test 4: File dengan ekstensi yang tidak cocok dengan isi aslinya ditolak."""
    fake_docx = tmp_path / "fake_document.docx"
    fake_docx.write_bytes(b"%PDF-1.5 Header palsu pada ekstensi docx")

    with pytest.raises(DocumentValidationError, match="Signature berkas tidak valid"):
        extractor.validate_file(fake_docx)


def test_oversized_file_validation(tmp_path: Path):
    """Test 5: File melebihi batas ukuran ditolak secara preventif."""
    small_extractor = UnifiedDocumentExtractor(max_size_mb=1)
    huge_file = tmp_path / "large_thesis.pdf"
    # Create file slightly above 1 MB
    with open(huge_file, "wb") as f:
        f.write(b"%PDF-1.4\n")
        f.seek(1024 * 1024 + 100)
        f.write(b"EOF")

    with pytest.raises(DocumentValidationError, match="melampaui batas maksimal"):
        small_extractor.validate_file(huge_file)


def test_ocr_engine_fallback_behavior():
    """Test 6: OCREngine menyediakan penanganan graceful ketika OCR tool atau bahasa tidak tersedia."""
    engine = OCREngine()
    readiness = engine.check_system_readiness()
    assert isinstance(readiness, dict)
    assert "tesseract_installed" in readiness
    assert "poppler_installed" in readiness

    # Test running OCR on a blank test image
    img = Image.new("RGB", (200, 100), color="white")
    draw = ImageDraw.Draw(img)
    draw.text((10, 10), "TESIS 2026", fill="black")

    res = engine.run_ocr(img)
    assert isinstance(res, str)


def test_document_converter_multi_format(tmp_path: Path):
    """Test 7: DocumentConverter mengubah dokumen ke format TXT, HTML, DOCX, dan PDF."""
    src_file = tmp_path / "source.txt"
    src_file.write_text("Bab 1 Pendahuluan\n\nPenelitian ini menguji pengaruh X terhadap Y.", encoding="utf-8")

    conv = DocumentConverter()

    out_html = tmp_path / "converted.html"
    conv.convert_document(src_file, "html", out_html)
    assert out_html.exists()
    assert "Bab 1 Pendahuluan" in out_html.read_text(encoding="utf-8")

    out_docx = tmp_path / "converted.docx"
    conv.convert_document(src_file, "docx", out_docx)
    assert out_docx.exists()
    doc = Document(str(out_docx))
    assert len(doc.paragraphs) >= 2

    out_pdf = tmp_path / "converted.pdf"
    conv.convert_document(src_file, "pdf", out_pdf)
    assert out_pdf.exists()
    assert out_pdf.stat().st_size > 500


def test_encrypted_pdf_rejection(tmp_path: Path):
    """Test 8: Dokumen PDF yang terenkripsi kata sandi ditolak dengan pesan yang jelas."""
    enc_pdf = tmp_path / "encrypted.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((50, 100), "Rahasia")
    # Save with user password
    doc.save(str(enc_pdf), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="kuncirahasia")
    doc.close()

    from engines.document_parser import parse_document
    with pytest.raises(ValueError, match="kata sandi|terenkripsi"):
        parse_document(str(enc_pdf))


def test_page_limit_175_exact_and_176_rejection(tmp_path: Path):
    """Test 9: Dokumen tepat 175 halaman diterima, sedangkan 176 halaman terdeteksi melebihi batas."""
    from engines.document_parser import parse_document

    # 1. Dokumen tepat 175 halaman
    p175 = tmp_path / "thesis_175_pages.pdf"
    doc175 = fitz.open()
    for _ in range(175):
        doc175.new_page().insert_text((50, 100), "Teks Halaman Tesis")
    doc175.save(str(p175))
    doc175.close()

    res175 = parse_document(str(p175))
    assert res175.get_effective_page_count() == 175
    assert not (res175.get_effective_page_count() > 175)

    # 2. Dokumen 176 halaman (melebihi batas)
    p176 = tmp_path / "thesis_176_pages.pdf"
    doc176 = fitz.open()
    for _ in range(176):
        doc176.new_page().insert_text((50, 100), "Teks Halaman Tesis")
    doc176.save(str(p176))
    doc176.close()

    res176 = parse_document(str(p176))
    assert res176.get_effective_page_count() == 176
    assert res176.get_effective_page_count() > 175


def test_docx_page_count_never_none(tmp_path: Path):
    """Test 10: Dokumen DOCX selalu menghasilkan get_effective_page_count() berupa bilangan bulat positif."""
    from engines.document_parser import parse_document

    sample_docx = tmp_path / "thesis_sample.docx"
    d = Document()
    d.add_heading("BAB I PENDAHULUAN", level=1)
    for _ in range(20):
        d.add_paragraph("Komunikasi organisasi memainkan peranan esensial dalam era disrupsi digital yang semakin masif.")
    d.save(str(sample_docx))

    res_docx = parse_document(str(sample_docx))
    effective_count = res_docx.get_effective_page_count()
    assert isinstance(effective_count, int)
    assert effective_count >= 1
    # Memastikan tidak pernah memicu TypeError saat dibandingkan dengan bilangan bulat
    assert not (effective_count > 175)


def test_boundary_word_counts_and_invalid_types():
    """Test 11: Uji kasus batas 0 kata, 1 kata, 48.125 kata, 48.126 kata, None, negatif, dan tipe invalid."""
    from models.document_model import ParsedDocument, ParsedParagraph

    # 1. 0 kata (dokumen kosong)
    doc_0 = ParsedDocument(filename="kosong.docx", file_type="docx", page_count=None)
    assert doc_0.get_effective_page_count() == 1
    assert doc_0.is_page_count_estimated is True

    # 2. 1 kata
    doc_1 = ParsedDocument(filename="satu.docx", file_type="docx", page_count=None)
    doc_1.paragraphs.append(ParsedParagraph(text="Pendahuluan"))
    assert doc_1.get_effective_page_count() == 1
    assert doc_1.is_page_count_estimated is True

    # 3. Tepat 48.125 kata -> 175 halaman (48.125 / 275 = 175.0)
    doc_175 = ParsedDocument(filename="batas_175.docx", file_type="docx", page_count=None)
    doc_175.paragraphs.append(ParsedParagraph(text="kata " * 48125))
    assert doc_175.get_effective_page_count() == 175
    assert not (doc_175.get_effective_page_count() > 175)
    assert doc_175.is_page_count_estimated is True

    # 4. 48.126 kata -> 176 halaman (ceil(48.126 / 275) = 176)
    doc_176 = ParsedDocument(filename="lebih_176.docx", file_type="docx", page_count=None)
    doc_176.paragraphs.append(ParsedParagraph(text="kata " * 48126))
    assert doc_176.get_effective_page_count() == 176
    assert doc_176.get_effective_page_count() > 175
    assert doc_176.is_page_count_estimated is True

    # 5. page_count = nilai negatif (-10)
    doc_neg = ParsedDocument(filename="negatif.docx", file_type="docx", page_count=-10)
    doc_neg.paragraphs.append(ParsedParagraph(text="kata " * 550))
    # Otomatis dikoreksi ke estimasi positif (550 / 275 = 2 halaman)
    assert doc_neg.get_effective_page_count() == 2
    assert doc_neg.is_page_count_estimated is True

    # 6. page_count = tipe tidak valid (string "invalid")
    doc_inv = ParsedDocument(filename="invalid.docx", file_type="docx", page_count="invalid")  # type: ignore
    doc_inv.paragraphs.append(ParsedParagraph(text="kata " * 275))
    assert doc_inv.get_effective_page_count() == 1
    assert doc_inv.is_page_count_estimated is True


def test_map_finding_to_evidence_handles_lists_and_cross_document_auditor():
    """Test 12: Memastikan map_finding_to_evidence tidak melempar AttributeError saat menerima list atau output tuple."""
    from engines.evidence_mapper import map_finding_to_evidence
    from engines.cross_document_auditor import CrossDocumentAuditor
    from models.document_model import ParsedDocument, ParsedParagraph

    doc_t = ParsedDocument(filename="tesis.docx", file_type="docx")
    doc_t.paragraphs.append(ParsedParagraph(text="Penelitian ini menggunakan pendekatan kuantitatif regresi."))

    doc_j = ParsedDocument(filename="jurnal.docx", file_type="docx")
    doc_j.paragraphs.append(ParsedParagraph(text="Artikel ini menggunakan pendekatan kualitatif naratif."))

    # CrossDocumentAuditor returns a 3-tuple: (findings, table, score)
    j_res = CrossDocumentAuditor(doc_t, doc_j).audit()
    assert isinstance(j_res, tuple)
    j_findings, j_table, j_score = j_res

    # Unpacked findings must map cleanly without error
    mapped = [map_finding_to_evidence(f, doc_t) for f in j_findings]
    assert len(mapped) == len(j_findings)

    # Defensive check: non-dict inputs (like an accidental list) must not crash
    raw_list_input = ["item_1", "item_2"]
    safe_mapped = map_finding_to_evidence(raw_list_input, doc_t)
    assert safe_mapped["rule_id"] == "GENERIC_FINDING"
    assert safe_mapped["mapping_status"] == "UNVERIFIED"

