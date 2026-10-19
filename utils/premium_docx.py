"""utils/premium_docx.py - Academic typography and Booktabs table styling for docx."""
import io
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def apply_premium_typography(doc: Document):
    """Menerapkan font Times New Roman 12pt, Justified, dan Widow/Orphan Control."""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Pengaturan paragraf
    paragraph_format = style.paragraph_format
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph_format.space_after = Pt(12)
    # Mengaktifkan Widow/Orphan control
    paragraph_format.widow_control = True


def set_booktabs_style(table):
    """Membuat tabel gaya Booktabs: Hanya garis horizontal, tanpa garis vertikal."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    
    # Hapus semua border bawaan
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        borders.append(border)
    tblPr.append(borders)

    # Tambahkan garis atas dan bawah tabel (tebal 1.5 pt / sz=12)
    top_border = OxmlElement('w:top')
    top_border.set(qn('w:val'), 'single')
    top_border.set(qn('w:sz'), '12')  # Ukuran garis
    borders.append(top_border)
    
    bottom_border = OxmlElement('w:bottom')
    bottom_border.set(qn('w:val'), 'single')
    bottom_border.set(qn('w:sz'), '12')
    borders.append(bottom_border)

    # Tambahkan garis bawah pada baris pertama (Header - 0.75 pt / sz=6)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        tcBorders.append(bottom)
        tcPr.append(tcBorders)


def generate_sample_premium_docx(
    title: str = "BAB I PENDAHULUAN",
    body_text: str = (
        "Penelitian ini menganalisis pengaruh transformasi komunikasi digital terhadap "
        "keterlibatan publik dalam program kebijakan gizi nasional. Seluruh naskah diformat menggunakan "
        "standar tipografi tesis MIKOM 2026 (Times New Roman 12pt, Rata Kanan-Kiri/Justified, Widow/Orphan Control aktif)."
    ),
    table_data: list = None
) -> io.BytesIO:
    """Menghasilkan dokumen docx siap unduh dengan format premium Times New Roman & tabel Booktabs."""
    doc = Document()
    apply_premium_typography(doc)
    
    doc.add_heading(title, level=1)
    doc.add_paragraph(body_text)
    
    if table_data is None:
        table_data = [
            ["Variabel", "Mean", "Std. Dev", "Sig."],
            ["Komunikasi Digital (X1)", "4.52", "0.68", "0.001"],
            ["Keterlibatan Publik (Y)", "4.38", "0.72", "0.003"],
            ["Kepercayaan Audiens (Z)", "4.61", "0.54", "0.002"]
        ]
        
    rows = len(table_data)
    cols = len(table_data[0]) if rows > 0 else 0
    if rows > 0 and cols > 0:
        table = doc.add_table(rows=rows, cols=cols)
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = str(val)
        set_booktabs_style(table)
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
