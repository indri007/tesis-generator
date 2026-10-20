"""Word counting utility for MIKOM 2026 thesis documents (Bab I–V excl. figures/tables/captions)."""

from __future__ import annotations

import re
from pathlib import Path
from docx import Document

WORD_MIN = 40_000
WORD_MAX = 60_000
MAX_PAGES = None


def count_words_excl_figures(path: str | Path | Any, start_pat: str = r"^BAB I\b", end_pat: str = r"^DAFTAR PUSTAKA") -> int:
    """Hitung kata isi Bab I–V, tanpa gambar, caption, dan tabel."""
    import io
    if hasattr(path, "read") or isinstance(path, (io.BytesIO, bytes)):
        try:
            doc = Document(path if not isinstance(path, bytes) else io.BytesIO(path))
            total, inside = 0, False
            for p in doc.paragraphs:
                t = p.text.strip()
                is_toc = (hasattr(p, "style") and p.style and hasattr(p.style, "name") and p.style.name and p.style.name.lower().startswith("toc")) or bool(re.search(r"\t\d+$", t))
                if is_toc:
                    continue
                is_heading = (hasattr(p, "style") and p.style and hasattr(p.style, "name") and p.style.name and p.style.name.startswith("Heading"))
                if re.match(start_pat, t, re.I) and (is_heading or re.match(r"^BAB\s+I(\b|:|\s)", t, re.I)):
                    inside = True
                if inside and re.match(end_pat, t, re.I):
                    break
                if not inside or not t:
                    continue
                if p._p.xpath(".//w:drawing | .//w:pict"):          # paragraf berisi gambar
                    continue
                if re.match(r"^(Gambar|Figure|Tabel|Table)\s*\d", t, re.I):  # caption
                    continue
                total += len(re.findall(r"\S+", t))
            return total
        except Exception:
            return 0

    file_path = Path(path)
    if not file_path.exists():
        return 0

    if file_path.suffix.lower() == ".docx":
        try:
            doc = Document(str(file_path))
            total, inside = 0, False
            for p in doc.paragraphs:
                t = p.text.strip()
                is_toc = (hasattr(p, "style") and p.style and hasattr(p.style, "name") and p.style.name and p.style.name.lower().startswith("toc")) or bool(re.search(r"\t\d+$", t))
                if is_toc:
                    continue
                is_heading = (hasattr(p, "style") and p.style and hasattr(p.style, "name") and p.style.name and p.style.name.startswith("Heading"))
                if re.match(start_pat, t, re.I) and (is_heading or re.match(r"^BAB\s+I(\b|:|\s)", t, re.I)):
                    inside = True
                if inside and re.match(end_pat, t, re.I):
                    break
                if not inside or not t:
                    continue
                if p._p.xpath(".//w:drawing | .//w:pict"):          # paragraf berisi gambar
                    continue
                if re.match(r"^(Gambar|Figure|Tabel|Table)\s*\d", t, re.I):  # caption
                    continue
                total += len(re.findall(r"\S+", t))
            return total
        except Exception:
            return 0

    elif file_path.suffix.lower() == ".pdf":
        try:
            import pymupdf
            doc_fitz = pymupdf.open(str(file_path))
            full_text = []
            for page in doc_fitz:
                full_text.append(page.get_text())
            doc_fitz.close()
            lines = "\n".join(full_text).splitlines()
            total, inside = 0, False
            for line in lines:
                t = line.strip()
                if re.match(start_pat, t, re.I):
                    inside = True
                if inside and re.match(end_pat, t, re.I):
                    break
                if not inside or not t:
                    continue
                if re.match(r"^(Gambar|Figure|Tabel|Table)\s*\d", t, re.I):
                    continue
                total += len(re.findall(r"\S+", t))
            return total
        except Exception:
            return 0

    return 0


def evaluate_word_count(words: int) -> tuple[str, str, str]:
    """Evaluates word count against 40,000–60,000 words target."""
    if words < WORD_MIN:
        return (
            "ERROR",
            f"Volume naskah ({words:,} kata) di bawah batas minimum 40.000 kata.",
            "Tambahkan elaborasi substantif pada Bab IV (Hasil dan Pembahasan) serta penguatan telaah pustaka Bab II."
        )
    elif words > WORD_MAX:
        return (
            "WARNING",
            f"Volume naskah ({words:,} kata) melebihi rentang rekomendasi 60.000 kata.",
            "Lakukan peringkasan pada bagian deskriptif agar naskah lebih padat dan fokus."
        )
    else:
        return (
            "PASS",
            f"Volume naskah ({words:,} kata) memenuhi standar MIKOM 2026 (40.000 – 60.000 kata).",
            "Pertahankan kedalaman pembahasan dan proporsi isi antar-bab."
        )
