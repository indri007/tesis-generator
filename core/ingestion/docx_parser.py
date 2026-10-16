from pathlib import Path

from docx import Document

from models.document_model import (
    ParsedDocument,
    ParsedParagraph,
    ParsedTable,
)


def _cm(value) -> float:
    return round(value / 360000, 2)


def parse_docx(path: str) -> ParsedDocument:
    file_path = Path(path)

    if not file_path.is_file():
        raise FileNotFoundError(path)

    doc = Document(str(file_path))

    result = ParsedDocument(
        filename=file_path.name,
        file_type="docx",
    )

    props = doc.core_properties
    result.metadata = {
        "title": props.title or "",
        "author": props.author or "",
        "subject": props.subject or "",
    }

    # Format section: nilai disimpan dalam cm.
    sections = []

    for section in doc.sections:
        sections.append({
            "page_width_cm": _cm(section.page_width),
            "page_height_cm": _cm(section.page_height),
            "left_margin_cm": _cm(section.left_margin),
            "right_margin_cm": _cm(section.right_margin),
            "top_margin_cm": _cm(section.top_margin),
            "bottom_margin_cm": _cm(section.bottom_margin),
        })

    result.format_data = {
        "sections": sections,
        # Untuk pemeriksaan banyak section, rule engine
        # perlu mengevaluasi setiap item sections.
        **(sections[0] if sections else {}),
    }

    # Urutan paragraf.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        result.paragraphs.append(
            ParsedParagraph(
                text=text,
                page=None,
                style=paragraph.style.name if paragraph.style else "",
            )
        )

    # Tabel.
    for table in doc.tables:
        rows = [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ]

        result.tables.append(ParsedTable(rows=rows))

    # Estimasi jumlah halaman berbasis densitas kata akademik (~275 kata/halaman)
    total_words = sum(len(p.text.split()) for p in result.paragraphs)
    result.page_count = max(1, round(total_words / 275)) if total_words > 0 else 1

    result.warnings.append(
        "DOCX: nomor halaman dan pemetaan posisi paragraf "
        "belum diverifikasi melalui rendering."
    )

    return result
