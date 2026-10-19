"""Revision Engine: Generates prioritized revision plans and automated formatting corrections."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


class RevisionEngine:
    """Classifies revision priorities and prepares corrected output files."""

    def __init__(
        self,
        findings: list[dict[str, Any]] | None = None,
        base_summary: dict[str, Any] | None = None,
        *args: Any,
        **kwargs: Any,
    ):
        self.findings = findings if findings is not None else []
        self.base_summary = base_summary or kwargs.get("base_summary", {})

    def calculate_scores(
        self,
        base_summary: dict[str, Any] | None = None,
        *args: Any,
        **kwargs: Any,
    ) -> dict[str, Any]:
        """Calculates compliance score and status counts strictly according to MIKOM 2026 rules.

        Formula:
            penalty = (ERROR * 15.0) + (WARNING * 5.0) + (UNVERIFIED * 2.0)
            overall_score = max(0.0, min(100.0, 100.0 - penalty))

        Returns:
            dict containing:
                - counts: dict with counts for PASS, ERROR, WARNING, INFO, UNVERIFIED
                - penalty: total deduction applied
                - overall_score: final compliance score in [0.0, 100.0]
                - category_scores: breakdown of scores per audit category
        """
        import math

        summary_source = base_summary if base_summary is not None else self.base_summary
        counts: dict[str, int] = {
            "PASS": int(summary_source.get("PASS", 0)) if isinstance(summary_source, dict) else 0,
            "ERROR": 0,
            "WARNING": 0,
            "INFO": 0,
            "UNVERIFIED": 0,
        }

        category_counts: dict[str, dict[str, int]] = {}

        for item in self.findings:
            if not isinstance(item, dict):
                counts["UNVERIFIED"] += 1
                continue

            raw_status = item.get("status") or item.get("finding_status") or item.get("severity") or "UNVERIFIED"
            status_norm = str(raw_status).strip().upper()

            if status_norm in ("ERROR", "ERR", "CRITICAL", "FATAL"):
                effective_status = "ERROR"
            elif status_norm in ("WARNING", "WARN"):
                effective_status = "WARNING"
            elif status_norm in ("PASS", "PASSED", "SUCCESS"):
                effective_status = "PASS"
            elif status_norm in ("INFO", "INFORMATION"):
                effective_status = "INFO"
            else:
                effective_status = "UNVERIFIED"

            counts[effective_status] += 1

            cat = str(item.get("category", "General")).strip() or "General"
            if cat not in category_counts:
                category_counts[cat] = {"ERROR": 0, "WARNING": 0, "UNVERIFIED": 0}
            if effective_status in category_counts[cat]:
                category_counts[cat][effective_status] += 1

        penalty = (counts["ERROR"] * 15.0) + (counts["WARNING"] * 5.0) + (counts["UNVERIFIED"] * 2.0)
        raw_score = 100.0 - penalty
        if math.isnan(raw_score) or math.isinf(raw_score):
            overall_score = 0.0
        else:
            overall_score = max(0.0, min(100.0, raw_score))

        category_scores: dict[str, float] = {}
        for cat, c_dict in category_counts.items():
            cat_pen = (c_dict["ERROR"] * 15.0) + (c_dict["WARNING"] * 5.0) + (c_dict["UNVERIFIED"] * 2.0)
            cat_score = max(0.0, min(100.0, 100.0 - cat_pen))
            category_scores[cat] = round(cat_score, 1)

        return {
            "counts": counts,
            "penalty": round(penalty, 1),
            "overall_score": round(overall_score, 1),
            "category_scores": category_scores,
        }

    def generate_checklist(self) -> dict[str, list[dict[str, Any]]]:
        """Categorizes findings into 3 actionable priority groups."""
        p1 = []  # Wajib Diperbaiki (Akademik, Rumus, Statistik, Sampel, Variabel)
        p2 = []  # Penting (Struktur bab, Sitasi, Rujukan)
        p3 = []  # Format (Margin, Font, Spasi)

        for f in self.findings:
            cat = f.get("category", "").lower()
            rule_id = f.get("rule_id", "")
            severity = f.get("status", "")

            if severity in ["PASS", "INFO"]:
                continue

            if any(k in cat for k in ["formula", "statistik", "variabel", "journal", "sample"]) or "RECALCULATION" in rule_id:
                p1.append(f)
            elif any(k in cat for k in ["struktur", "sitasi", "referensi", "abstrak"]):
                p2.append(f)
            else:
                p3.append(f)

        return {
            "PRIORITAS 1 — WAJIB DIPERBAIKI (Substantif & Validitas Ilmiah)": p1,
            "PRIORITAS 2 — PENTING (Struktur & Kelengkapan Akademik)": p2,
            "PRIORITAS 3 — FORMAT & TATA LETAK (Mekanik)": p3,
        }

    def apply_docx_formatting_fixes(self, input_path: str, output_path: str) -> bool:
        """Applies safe mechanical formatting corrections to a DOCX document:
        - Margins: Top 4cm, Left 4cm, Bottom 3cm, Right 3cm
        - Fonts: Times New Roman, 12pt
        - Spacing: Double (2.0)
        - Alignment: Justified
        """
        try:
            doc = Document(input_path)
            for section in doc.sections:
                section.top_margin = Cm(4.0)
                section.left_margin = Cm(4.0)
                section.bottom_margin = Cm(3.0)
                section.right_margin = Cm(3.0)

            for p in doc.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 2.0
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)

            doc.save(output_path)
            return True
        except Exception:
            return False


class DocxAutoFixer:
    """Automates mechanical formatting corrections according to MIKOM 2026 standards.
    
    Standards enforced:
    - Page Margins: Top 4.0 cm, Left 4.0 cm, Bottom 3.0 cm, Right 3.0 cm.
    - Typography: Times New Roman, 12 pt for body text.
    - Spacing: Double (2.0) line spacing for academic body paragraphs.
    - Alignment: Justified (Rata Kiri-Kanan) while preserving centered headings.
    - Table formatting: Cleans and preserves table cell text.
    """

    def __init__(self, docx_path: str):
        self.docx_path = Path(docx_path)
        if not self.docx_path.exists():
            raise FileNotFoundError(f"Berkas DOCX sumber tidak ditemukan: {docx_path}")

    def fix_all(self, output_path: str) -> list[dict[str, Any]]:
        """Applies 4-4-3-3 cm margins, Times New Roman 12pt, double spacing, and justification.
        
        Raises:
            ValueError: If the DOCX file is corrupt or cannot be parsed.
            IOError: If saving to the destination path fails.
        """
        try:
            doc = Document(str(self.docx_path))
        except Exception as e:
            raise ValueError(f"Berkas DOCX rusak atau tidak dapat diuraikan: {e}") from e

        log: list[dict[str, Any]] = []

        # 1. Apply Margins: 4.0 cm (Top, Left) and 3.0 cm (Bottom, Right)
        section_count = 0
        for section in doc.sections:
            section.top_margin = Cm(4.0)
            section.left_margin = Cm(4.0)
            section.bottom_margin = Cm(3.0)
            section.right_margin = Cm(3.0)
            section_count += 1
        log.append({
            "action": "SET_MARGINS_4_4_3_3_CM",
            "sections_affected": section_count,
            "status": "APPLIED",
            "detail": "Margin disesuaikan ke Kiri: 4cm, Atas: 4cm, Kanan: 3cm, Bawah: 3cm",
        })

        # 2. Format Paragraphs & Preserve Headings
        body_p_count = 0
        heading_count = 0
        for p in doc.paragraphs:
            style_name = p.style.name.lower() if p.style and p.style.name else ""
            is_heading = "heading" in style_name or "judul" in style_name or p.alignment == WD_ALIGN_PARAGRAPH.CENTER

            if is_heading:
                heading_count += 1
                # Preserve alignment for headings, but ensure Times New Roman font
                for r in p.runs:
                    r.font.name = "Times New Roman"
            else:
                body_p_count += 1
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 2.0
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)

        log.append({
            "action": "SET_FONT_AND_SPACING",
            "body_paragraphs": body_p_count,
            "headings_preserved": heading_count,
            "status": "APPLIED",
            "detail": "Font Times New Roman 12pt, spasi 2.0, dan perataan justify berhasil diterapkan.",
        })

        # 3. Format Tables safely without corrupting XML
        table_cells_count = 0
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_cells_count += 1
                    for cp in cell.paragraphs:
                        for cr in cp.runs:
                            cr.font.name = "Times New Roman"

        if doc.tables:
            log.append({
                "action": "FORMAT_TABLE_CELLS",
                "tables_count": len(doc.tables),
                "cells_affected": table_cells_count,
                "status": "APPLIED",
                "detail": "Tipografi sel tabel disinkronkan ke Times New Roman.",
            })

        # 4. Save to output path safely
        out_p = Path(output_path)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        try:
            doc.save(str(out_p))
        except Exception as e:
            raise IOError(f"Gagal menyimpan dokumen DOCX hasil perbaikan ke {output_path}: {e}") from e

        return log


