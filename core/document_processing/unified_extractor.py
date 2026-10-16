"""Unified Document Processing & Extraction Engine for PDF, DOCX, and Images with OCR fallback."""

import logging
from pathlib import Path
from typing import Any, Optional

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from docx import Document
from PIL import Image

from core.ocr.ocr_engine import OCREngine

logger = logging.getLogger(__name__)

# Magic byte signatures
MAGIC_SIGNATURES = {
    ".pdf": b"%PDF-",
    ".docx": b"PK\x03\x04",
    ".png": b"\x89PNG\r\n\x1a\n",
    ".jpg": b"\xff\xd8\xff",
    ".jpeg": b"\xff\xd8\xff",
}


class DocumentValidationError(ValueError):
    """Raised when an uploaded document fails security or format validation."""
    pass


class UnifiedDocumentExtractor:
    """Provides high-level, secure, and validated document extraction."""

    def __init__(self, max_size_mb: int = 50, ocr_engine: Optional[OCREngine] = None):
        self.max_size_mb = max_size_mb
        self.max_bytes = max_size_mb * 1024 * 1024
        self.ocr_engine = ocr_engine or OCREngine()

    def validate_file(self, file_path: str | Path) -> dict[str, Any]:
        """Validates file integrity, extension match, size limits, and encryption status."""
        p = Path(file_path)
        if not p.exists():
            raise DocumentValidationError(f"Berkas tidak ditemukan: {p.name}")

        file_size = p.stat().st_size
        if file_size == 0:
            raise DocumentValidationError(f"Berkas kosong (0 byte): {p.name}")
        if file_size > self.max_bytes:
            raise DocumentValidationError(
                f"Ukuran berkas ({file_size / (1024 * 1024):.1f} MB) melampaui batas maksimal {self.max_size_mb} MB"
            )

        suffix = p.suffix.lower()
        if suffix not in MAGIC_SIGNATURES and suffix not in [".txt"]:
            raise DocumentValidationError(f"Format ekstensi '{suffix}' tidak didukung.")

        # Validate Magic Bytes
        if suffix in MAGIC_SIGNATURES:
            expected_magic = MAGIC_SIGNATURES[suffix]
            with open(p, "rb") as f:
                header = f.read(len(expected_magic))
                if not header.startswith(expected_magic):
                    raise DocumentValidationError(
                        f"Signature berkas tidak valid untuk format {suffix}. Berkas mungkin rusak atau berekstensi palsu."
                    )

        # Check PDF Encryption
        is_encrypted = False
        page_count = 0
        if suffix == ".pdf":
            try:
                doc = fitz.open(str(p))
                is_encrypted = doc.is_encrypted
                page_count = doc.page_count
                doc.close()
                if is_encrypted:
                    raise DocumentValidationError("Dokumen PDF terenkripsi/diproteksi kata sandi. Harap buka kunci terlebih dahulu.")
            except Exception as e:
                if isinstance(e, DocumentValidationError):
                    raise
                raise DocumentValidationError(f"Dokumen PDF rusak atau tidak dapat dibaca: {e}")

        return {
            "valid": True,
            "filename": p.name,
            "suffix": suffix,
            "size_bytes": file_size,
            "is_encrypted": is_encrypted,
            "page_count": page_count,
        }

    def extract_pdf_text(self, file_path: str | Path, fallback_to_ocr: bool = True) -> str:
        """Extracts text from PDF page by page, with automated OCR fallback for scanned pages."""
        self.validate_file(file_path)
        p = Path(file_path)
        doc = fitz.open(str(p))
        pages_content = []

        try:
            for i, page in enumerate(doc):
                text = page.get_text().strip()
                # If page text is very sparse (< 30 chars), attempt OCR on that page if requested
                if len(text) < 30 and fallback_to_ocr and self.ocr_engine.is_tesseract_installed:
                    logger.info("Page %d has insufficient digital text. Falling back to OCR.", i + 1)
                    pix = page.get_pixmap(dpi=150)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr_text = self.ocr_engine.run_ocr(img)
                    if len(ocr_text) > len(text):
                        text = f"[Hasil OCR Halaman {i + 1}]\n" + ocr_text

                pages_content.append(text)
        finally:
            doc.close()

        return "\n\n".join(pages_content)

    def extract_docx_text(self, file_path: str | Path) -> str:
        """Extracts all text from Word DOCX including paragraphs and tables."""
        self.validate_file(file_path)
        doc = Document(str(file_path))
        lines = []

        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    lines.append(" | ".join(row_text))

        return "\n\n".join(lines)

    def extract_text(self, file_path: str | Path) -> str:
        """Generic entry point to extract text across supported document formats."""
        p = Path(file_path)
        suffix = p.suffix.lower()

        if suffix == ".pdf":
            return self.extract_pdf_text(p)
        elif suffix == ".docx":
            return self.extract_docx_text(p)
        elif suffix in [".png", ".jpg", ".jpeg"]:
            self.validate_file(p)
            with Image.open(str(p)) as img:
                return self.ocr_engine.run_ocr(img)
        elif suffix == ".txt":
            self.validate_file(p)
            return p.read_text(encoding="utf-8", errors="replace")
        else:
            raise DocumentValidationError(f"Format berkas '{suffix}' tidak didukung untuk ekstraksi teks.")
