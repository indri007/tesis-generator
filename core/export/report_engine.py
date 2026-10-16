"""Report Engine: Generates multi-format audit reports (HTML, PDF, DOCX, JSON)."""

from __future__ import annotations

import io
import json
from typing import Any

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class ReportEngine:
    """Generates official audit reports in various formats."""

    def __init__(
        self,
        summary: dict[str, Any],
        findings: list[dict[str, Any]],
        document_info: dict[str, Any],
        checklist: dict[str, list[dict[str, Any]]],
        category_scores: dict[str, float] | None = None,
        overall_score: float = 85.0
    ):
        self.summary = summary
        self.findings = findings
        self.doc_info = document_info
        self.checklist = checklist
        self.category_scores = category_scores or {}
        self.overall_score = overall_score

    def to_json(self) -> str:
        """Exports audit report as structured JSON."""
        data = {
            "title": "LAPORAN AUDIT KEPATUHAN & KONSISTENSI TESIS MIKOM 2026",
            "document": self.doc_info,
            "overall_score": self.overall_score,
            "category_scores": self.category_scores,
            "summary": self.summary,
            "findings_count": len(self.findings),
            "findings": self.findings,
            "revision_checklist": self.checklist
        }
        return json.dumps(data, indent=2, ensure_ascii=False)

    def to_html(self) -> str:
        """Exports executive HTML audit report."""
        rows = ""
        for f in self.findings:
            status_color = "#DC2626" if f.get("status") == "ERROR" else "#D97706"
            rows += f"""
            <tr style="border-bottom: 1px solid #E2E8F0;">
                <td style="padding: 10px; font-weight: bold; color: {status_color};">{f.get('status')}</td>
                <td style="padding: 10px; font-family: monospace;">{f.get('rule_id')}</td>
                <td style="padding: 10px;">{f.get('category')}</td>
                <td style="padding: 10px;"><b>Hal. {f.get('page') or '-'}</b>: {f.get('description')}</td>
                <td style="padding: 10px; font-size: 0.9em; background: #F8FAFC;">{f.get('actual')}</td>
                <td style="padding: 10px; font-size: 0.9em;">{f.get('recommendation')}</td>
            </tr>
            """

        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>Laporan Audit Tesis MIKOM 2026</title>
            <style>
                body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; margin: 40px; color: #1E293B; }}
                h1 {{ color: #0F172A; border-bottom: 2px solid #3B82F6; padding-bottom: 10px; }}
                .badge {{ padding: 6px 12px; border-radius: 6px; font-weight: bold; color: white; display: inline-block; }}
                .score {{ font-size: 32px; font-weight: bold; color: #2563EB; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 20px; }}
                th {{ background: #F1F5F9; padding: 12px; text-align: left; font-size: 0.85em; text-transform: uppercase; }}
            </style>
        </head>
        <body>
            <h1>Laporan Audit Kepatuhan & Konsistensi Tesis MIKOM 2026</h1>
            <p><b>Dokumen:</b> {self.doc_info.get('filename')} | <b>Halaman:</b> {self.doc_info.get('page_count')} | <b>Total Temuan:</b> {len(self.findings)}</p>
            <div style="background: #EFF6FF; border: 1px solid #BFDBFE; border-radius: 8px; padding: 16px; margin: 20px 0;">
                <span class="score">{self.overall_score} / 100</span>
                <p style="margin: 5px 0 0 0; color: #1E40AF;">Indeks Kepatuhan & Konsistensi Ilmiah Tesis</p>
            </div>
            <table>
                <thead>
                    <tr>
                        <th>Status</th>
                        <th>ID Aturan</th>
                        <th>Kategori</th>
                        <th>Deskripsi & Lokasi</th>
                        <th>Kondisi Terdeteksi</th>
                        <th>Rekomendasi Revisi</th>
                    </tr>
                </thead>
                <tbody>
                    {rows}
                </tbody>
            </table>
        </body>
        </html>
        """
        return html

    def to_docx_bytes(self) -> io.BytesIO:
        """Exports audit report as a DOCX document."""
        doc = Document()
        doc.add_heading("LAPORAN AUDIT TESIS MIKOM 2026", level=1)
        doc.add_paragraph(f"Dokumen: {self.doc_info.get('filename')} | Skor: {self.overall_score}/100")

        # Summary counts
        p = doc.add_paragraph()
        p.add_run(f"PASS: {self.summary.get('PASS', 0)}  |  ERROR: {self.summary.get('ERROR', 0)}  |  WARNING: {self.summary.get('WARNING', 0)}  |  UNVERIFIED: {self.summary.get('UNVERIFIED', 0)}").bold = True

        doc.add_heading("Rincian Temuan Audit", level=2)
        for f in self.findings:
            item_p = doc.add_paragraph()
            run_status = item_p.add_run(f"[{f.get('status')}] {f.get('rule_id')}: {f.get('description')}\n")
            run_status.bold = True
            item_p.add_run(f"Lokasi: Halaman {f.get('page') or '-'} | Bagian: {f.get('section') or '-'}\n")
            item_p.add_run(f"Aktual: {f.get('actual')}\n")
            item_p.add_run(f"Ekspektasi: {f.get('expected')}\n")
            item_p.add_run(f"Rekomendasi: {f.get('recommendation')}\n")

        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)
        return bio

    def to_pdf_bytes(self) -> io.BytesIO:
        """Exports audit report as PDF using ReportLab."""
        bio = io.BytesIO()
        pdf = SimpleDocTemplate(bio, pagesize=A4)
        styles = getSampleStyleSheet()

        elements = []
        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            textColor=colors.HexColor("#1E3A8A"),
            spaceAfter=12
        )
        elements.append(Paragraph("LAPORAN AUDIT TESIS MIKOM 2026", title_style))
        elements.append(Paragraph(f"Dokumen: {self.doc_info.get('filename')} | Skor Kepatuhan: {self.overall_score}/100", styles["Normal"]))
        elements.append(Spacer(1, 15))

        # Table of findings
        table_data = [["Status", "ID Aturan", "Halaman", "Deskripsi", "Rekomendasi"]]
        for f in self.findings[:25]:  # Top 25 findings
            table_data.append([
                str(f.get("status", "")),
                str(f.get("rule_id", "")),
                str(f.get("page") or "-"),
                str(f.get("description", ""))[:45],
                str(f.get("recommendation", ""))[:45]
            ])

        t = Table(table_data, colWidths=[50, 75, 45, 170, 170])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elements.append(t)
        pdf.build(elements)
        bio.seek(0)
        return bio
