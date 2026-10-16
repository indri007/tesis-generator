"""Document conversion module for transforming academic files across formats."""

from pathlib import Path
from typing import Union
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from core.document_processing.unified_extractor import UnifiedDocumentExtractor


class DocumentConverter:
    """Handles multi-format academic document conversions."""

    def __init__(self):
        self.extractor = UnifiedDocumentExtractor()

    def convert_document(self, input_file: str | Path, output_format: str, output_path: str | Path) -> Path:
        """Converts an input document (PDF, DOCX, TXT) into the specified output format."""
        in_p = Path(input_file)
        out_p = Path(output_path)
        out_fmt = output_format.lower().lstrip(".")

        text = self.extractor.extract_text(in_p)

        if out_fmt == "txt":
            out_p.write_text(text, encoding="utf-8")
        elif out_fmt == "html":
            paragraphs = "".join(f"<p>{p.strip()}</p>\n" for p in text.split("\n\n") if p.strip())
            html_content = (
                f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{in_p.name}</title></head>"
                f"<body style='font-family: Arial, sans-serif; line-height: 1.6; max-width: 800px; margin: 40px auto;'>"
                f"<h2>Hasil Konversi: {in_p.name}</h2>{paragraphs}</body></html>"
            )
            out_p.write_text(html_content, encoding="utf-8")
        elif out_fmt == "docx":
            doc = Document()
            doc.add_heading(f"Konversi Dokumen: {in_p.name}", level=1)
            for paragraph in text.split("\n\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph.strip())
            doc.save(str(out_p))
        elif out_fmt == "pdf":
            c = canvas.Canvas(str(out_p), pagesize=A4)
            width, height = A4
            y = height - 50
            c.setFont("Helvetica-Bold", 14)
            c.drawString(50, y, f"Konversi Dokumen: {in_p.name}")
            y -= 30
            c.setFont("Helvetica", 10)

            for line in text.splitlines():
                if y < 50:
                    c.showPage()
                    c.setFont("Helvetica", 10)
                    y = height - 50
                c.drawString(50, y, line[:90])
                y -= 14
            c.save()
        else:
            raise ValueError(f"Format output '{output_format}' tidak didukung. Pilihan: txt, html, docx, pdf.")

        return out_p
